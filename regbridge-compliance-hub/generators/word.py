"""Générateurs Word : trames d'examen renforcé, fiches KYC, procédures, DS, PPE, gel, politiques…"""
from __future__ import annotations

import io
from datetime import date

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
TEAL = RGBColor(0x0E, 0x74, 0x90)
GREY = RGBColor(0x5B, 0x6B, 0x82)
BOX = "☐"


# ------------------------------------------------------------------ socle
def _shade(cell, hex_color: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def new_doc(title: str, company: str, subtitle: str = "", confid: str = "Document interne — Confidentiel") -> Document:
    d = Document()
    sec = d.sections[0]
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = Cm(1.8)
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    for lvl, size, color in ((1, 15, NAVY), (2, 12.5, TEAL), (3, 11, NAVY)):
        h = d.styles[f"Heading {lvl}"]
        h.font.name, h.font.size, h.font.color.rgb, h.font.bold = "Calibri", Pt(size), color, True
    hp = sec.header.paragraphs[0]
    hp.text = f"{company or 'Entité'}  ·  {title}"
    hp.runs[0].font.size, hp.runs[0].font.color.rgb = Pt(8.5), GREY
    fp = sec.footer.paragraphs[0]
    fp.text = f"{confid} — généré par RegBridge Compliance Hub le {date.today():%d/%m/%Y}"
    fp.runs[0].font.size, fp.runs[0].font.color.rgb = Pt(8), GREY
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = d.add_paragraph()
    r = p.add_run(title.upper())
    r.bold, r.font.size, r.font.color.rgb = True, Pt(19), NAVY
    if subtitle:
        p2 = d.add_paragraph()
        r2 = p2.add_run(subtitle)
        r2.italic, r2.font.size, r2.font.color.rgb = True, Pt(10.5), TEAL
    return d


def table(d: Document, rows: list[list], header: list[str] | None = None, widths: list[float] | None = None, label_col=False):
    n = len(header or rows[0])
    t = d.add_table(rows=0, cols=n)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    if header:
        cells = t.add_row().cells
        for i, h in enumerate(header):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(h))
            run.bold, run.font.color.rgb, run.font.size = True, RGBColor(0xFF, 0xFF, 0xFF), Pt(9.5)
            _shade(cells[i], "0B1F3A")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = "" if v is None else str(v)
            for p in cells[i].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9.5)
            if label_col and i == 0:
                _shade(cells[i], "E6F4F7")
                for r in cells[i].paragraphs[0].runs:
                    r.bold = True
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    d.add_paragraph()
    return t


def fields(d: Document, labels: list[str]):
    table(d, [[l, ""] for l in labels], widths=[6.5, 10.5], label_col=True)


def checklist(d: Document, items: list[str]):
    for it in items:
        d.add_paragraph(f"{BOX}  {it}")


def signatures(d: Document, roles: list[str]):
    d.add_heading("Validation", 2)
    table(d, [[r, "", "", ""] for r in roles], header=["Fonction", "Nom", "Date", "Signature"], widths=[5, 4.5, 3, 4.5])


def to_bytes(d: Document) -> bytes:
    bio = io.BytesIO()
    d.save(bio)
    return bio.getvalue()


# ------------------------------------------------------------------ trames
def edd_template(client_type: str, sector_label: str, company: str, typologies: pd.DataFrame, kyc: pd.DataFrame,
                 context: dict | None = None) -> bytes:
    """Trame d'examen renforcé (CMF L.561-10-2 ; AMLR art. 34)."""
    ctx = context or {}
    d = new_doc("Trame d'examen renforcé", company, f"{client_type} — secteur {sector_label} — CMF L.561-10-2 / AMLR art. 34")
    d.add_paragraph("Objet : documenter l'analyse d'une opération particulièrement complexe, d'un montant inhabituellement élevé "
                    "ou ne paraissant pas avoir de justification économique ou d'objet licite. L'examen doit être finalisé dans "
                    "les délais fixés par la procédure interne (cf. décision ACPR Banque Delubac du 19/06/2025 : 66 examens renforcés "
                    "en attente depuis plus de 120 jours ont été sanctionnés).")
    d.add_heading("1. Déclenchement", 1)
    fields(d, ["N° de dossier", "Date de déclenchement", "Origine (alerte automatique / signalement / réquisition / presse)",
               "Scénario ou motif", "Analyste en charge", "Date limite de finalisation", ctx.get("extra_label", "Montant(s) concerné(s)")])
    d.add_heading("2. Connaissance du client", 1)
    if "physique" in client_type.lower() or "PP" in client_type:
        fields(d, ["Nom, prénom, date et lieu de naissance", "Nationalité(s) / résidence fiscale", "Profession / employeur",
                   "Revenus annuels déclarés", "Patrimoine estimé", "Statut PPE / proche de PPE", "Ancienneté de la relation",
                   "Niveau de risque actuel", "Fonctionnement attendu du compte (déclaré à l'entrée en relation)"])
    else:
        fields(d, ["Dénomination, forme, n° d'immatriculation", "Pays d'immatriculation / d'activité", "Objet social et activité réelle",
                   "Chiffre d'affaires / total bilan", "Bénéficiaires effectifs (≥ 25 %) et organigramme", "Dirigeants et statut PPE",
                   "Principaux clients / fournisseurs / pays de flux", "Niveau de risque actuel", "Fonctionnement attendu (flux prévisionnels)"])
    d.add_heading("3. Analyse des opérations", 1)
    table(d, [["", "", "", "", "", ""] for _ in range(5)],
          header=["Date", "Nature", "Montant", "Contrepartie", "Pays", "Cohérence avec le profil ?"], widths=[2.2, 3, 2.5, 3.5, 2, 3.8])
    fields(d, ["Origine des fonds (justificatifs)", "Destination des fonds", "Justification économique fournie par le client",
               "Cohérence avec revenus / patrimoine / CA déclarés", "Fractionnement, rebonds, comptes de passage ?",
               "Liens avec pays GAFI / liste UE / sanctions (opérations effectives)"])
    d.add_heading("4. Recherches effectuées", 1)
    checklist(d, ["Filtrage listes de gel UE / ONU / DG Trésor (et OFAC si exposition USD)", "Filtrage PPE (titulaire, BE, mandataires)",
                  "Recherche presse négative (FR/EN) — corruption, fraude, blanchiment, terrorisme", "Consultation RBE / registres du commerce",
                  "Historique des DS et réquisitions judiciaires sur le client et les personnes liées",
                  "Analyse des comptes liés et des personnes liées (famille, sociétés)", "Échange avec le chargé de clientèle"])
    d.add_heading("5. Typologies et critères d'alerte pertinents", 1)
    if len(typologies):
        table(d, [[r.intitule, r.criteres_alerte] for r in typologies.head(8).itertuples()], header=["Typologie", "Critères d'alerte (Tracfin)"],
              widths=[5.5, 11.5])
    d.add_heading("6. Pièces justificatives demandées", 1)
    for r in kyc.itertuples():
        d.add_paragraph(f"{BOX}  {r.piece}  ({r.reference})")
    d.add_heading("7. Conclusion", 1)
    checklist(d, ["Opération justifiée — classement documenté", "Maintien / mise en vigilance renforcée et révision du profil de risque",
                  "Déclaration de soupçon à Tracfin (Ermes) — préalable à l'exécution si possible (CMF L.561-15)",
                  "Fin de la relation d'affaires — à apprécier AVEC examen de l'opportunité d'une DS (cf. décision Delubac, grief 2)",
                  "Information du responsable LCB-FT / comité"])
    fields(d, ["Motivation détaillée de la conclusion", "Date de finalisation", "Délai de traitement (jours)"])
    signatures(d, ["Analyste LCB-FT", "Responsable LCB-FT", "Direction (si PPE / DS)"])
    return to_bytes(d)


def kyc_form(client_type: str, company: str, kyc: pd.DataFrame) -> bytes:
    d = new_doc("Fiche de recueil KYC", company, f"Typologie : {client_type} — CMF L.561-5 et s. / AMLR art. 19-22")
    d.add_heading("1. Identification", 1)
    if any(k in client_type.lower() for k in ("physique", "mineur", "entrepreneur", "ppe")):
        fields(d, ["Civilité, nom, prénom(s)", "Nom de naissance", "Date et lieu de naissance", "Nationalité(s)", "Adresse complète",
                   "Résidence fiscale / NIF", "Téléphone / e-mail", "Pièce d'identité (type, n°, validité)"])
        d.add_heading("2. Situation professionnelle et financière", 1)
        fields(d, ["Profession / secteur", "Employeur", "Revenus annuels nets", "Patrimoine estimé (fourchette)", "Origine du patrimoine",
                   "Statut PPE (titulaire / famille / proche)"])
    else:
        fields(d, ["Dénomination sociale / nom commercial", "Forme juridique", "N° SIREN / LEI / immatriculation", "Siège social",
                   "Pays d'activité", "Date de création", "Représentant légal (identité, pouvoirs)", "Code NAF / activité réelle"])
        d.add_heading("2. Bénéficiaires effectifs", 1)
        table(d, [["", "", "", "", ""] for _ in range(4)], header=["Nom", "Date de naissance", "% capital / droits de vote", "Contrôle (autre)", "PPE ?"],
              widths=[4.5, 3, 3.5, 3.5, 2.5])
        fields(d, ["Chiffre d'affaires / total bilan", "Effectif", "Principaux clients et fournisseurs", "Pays des flux attendus"])
    d.add_heading("3. Relation d'affaires", 1)
    fields(d, ["Objet et nature de la relation", "Produits / services souscrits", "Canal d'entrée en relation (face-à-face / distance / intermédiaire)",
               "Flux mensuels attendus (montant, nombre)", "Origine des fonds initiaux", "Recours aux espèces (oui/non, volume)"])
    d.add_heading("4. Pièces justificatives", 1)
    for lvl in ["Simplifiée", "Standard", "Renforcée"]:
        sub = kyc[kyc["niveau_vigilance"] == lvl]
        if len(sub):
            d.add_heading(f"Vigilance {lvl.lower()}", 3)
            for r in sub.itertuples():
                d.add_paragraph(f"{BOX}  {r.piece}  — {r.reference}")
    d.add_heading("5. Classification du risque", 1)
    table(d, [["Pays", "", ""], ["Produit / service", "", ""], ["Canal de distribution", "", ""], ["Client (PPE, structure, activité)", "", ""],
              ["Score global", "", ""]], header=["Critère", "Cotation (1-4)", "Commentaire"], widths=[6, 3, 8], label_col=True)
    checklist(d, ["Filtrage sanctions/gel effectué (date et résultat)", "Filtrage PPE effectué", "Recherche presse négative effectuée"])
    signatures(d, ["Chargé de clientèle", "Contrôle KYC (4 yeux)", "Responsable LCB-FT (si risque élevé)"])
    return to_bytes(d)


def ds_template(company: str, sector_label: str) -> bytes:
    d = new_doc("Trame de déclaration de soupçon", company, f"À transmettre via Ermes (Tracfin) — secteur {sector_label} — CMF L.561-15")
    d.add_paragraph("Rappel : la DS est confidentielle — interdiction d'en informer le client (CMF L.561-18). Elle doit en principe être "
                    "effectuée avant l'exécution de l'opération. Conserver le dossier 5 ans.")
    for h, items in [
        ("1. Déclarant", ["Organisme / n° d'agrément", "Correspondant / déclarant Tracfin", "Date"]),
        ("2. Personnes concernées", ["Identité complète (PP/PM)", "Bénéficiaires effectifs", "Personnes liées", "Statut PPE / presse négative"]),
        ("3. Opérations", ["Liste des opérations (date, montant, contrepartie, pays)", "Comptes concernés (IBAN)", "Opérations à venir / en attente"]),
        ("4. Analyse", ["Circonstances de détection", "Éléments d'atypisme vs profil client", "Recherches réalisées",
                        "Typologie suspectée (cf. Tracfin)", "Infraction sous-jacente présumée"]),
        ("5. Motifs du soupçon", ["Synthèse argumentée", "Pièces jointes (relevés, justificatifs, captures)"]),
    ]:
        d.add_heading(h, 1)
        fields(d, items)
    checklist(d, ["Exécution de l'opération suspendue ?", "Droit d'opposition de Tracfin envisageable (CMF L.561-24)",
                  "Mise à jour du profil de risque et vigilance renforcée", "Inscription au registre des DS"])
    signatures(d, ["Déclarant Tracfin", "Correspondant Tracfin"])
    return to_bytes(d)


def pep_form(company: str) -> bytes:
    d = new_doc("Fiche d'approbation PPE", company, "Personne politiquement exposée — CMF L.561-10 2° / R.561-20-2 ; AMLR art. 42-46")
    fields(d, ["Nom du client / BE", "Fonction exercée (cf. liste JOUE C/2023/724)", "Pays", "Catégorie (titulaire / famille / proche associé)",
               "Date de début / fin de fonction", "Niveau de risque pays (IPC, GAFI)"])
    d.add_heading("Origine du patrimoine et des fonds", 1)
    fields(d, ["Patrimoine estimé", "Origine du patrimoine (justificatifs)", "Origine des fonds de la relation", "Cohérence avec rémunération publique"])
    d.add_heading("Mesures de vigilance renforcée", 1)
    checklist(d, ["Accord d'un membre de la direction", "Revue au moins annuelle", "Surveillance renforcée des opérations (seuils dédiés)",
                  "Presse négative vérifiée (corruption, détournement)", "Maintien du statut ≥ 12 mois après fin de fonction"])
    signatures(d, ["Responsable LCB-FT", "Membre de la direction (accord)"])
    return to_bytes(d)


def sanctions_alert_form(company: str) -> bytes:
    d = new_doc("Fiche de traitement d'une alerte de filtrage", company, "Gel des avoirs et sanctions — CMF L.562-4-1 ; arrêté du 6 janvier 2021")
    fields(d, ["N° d'alerte / date", "Nom filtré (client, BE, contrepartie)", "Liste concernée (UE / ONU / FR / OFAC / UK)", "Nom de la liste et programme",
               "Score de correspondance"])
    d.add_heading("Analyse discriminante", 1)
    table(d, [["Nom / alias", "", "", ""], ["Date de naissance", "", "", ""], ["Nationalité", "", "", ""], ["Adresse / pays", "", "", ""],
              ["Identifiants (passeport, n° société)", "", "", ""]], header=["Critère", "Client", "Liste", "Concordance ?"], widths=[5, 4, 4, 4], label_col=True)
    checklist(d, ["Faux positif — justification documentée", "Homonymie non levée — escalade", "Vrai positif — gel immédiat des fonds",
                  "Déclaration à la DG Trésor (gel)", "Information Tracfin si soupçon", "Contournement suspecté (Règl. 833/2014 art. 12)"])
    signatures(d, ["Analyste", "Responsable gel / LCB-FT"])
    return to_bytes(d)


def procedure(company: str, sector_code: str, sector_label: str, factors: pd.DataFrame, typologies: pd.DataFrame, regs: pd.DataFrame) -> bytes:
    d = new_doc("Procédure LCB-FT et gel des avoirs", company, f"Secteur : {sector_label} — Version 1.0")
    d.add_heading("1. Objet et champ d'application", 1)
    d.add_paragraph(f"La présente procédure décrit le dispositif de lutte contre le blanchiment de capitaux, le financement du terrorisme et "
                    f"de mise en œuvre des mesures de gel des avoirs de {company or 'l’entité'}. Elle s'applique à l'ensemble des collaborateurs, "
                    "prestataires et intermédiaires. Elle anticipe l'application du règlement (UE) 2024/1624 (AMLR) au 10 juillet 2027.")
    d.add_heading("2. Cadre réglementaire", 1)
    for r in regs.head(10).itertuples():
        d.add_paragraph(f"{r._2} — {r.Statut} ({r._5})", style="List Bullet")
    d.add_heading("3. Gouvernance et responsabilités", 1)
    table(d, [["Organe de direction", "Approuve la classification des risques et la politique ; membre responsable LCB-FT (AMLR art. 11)"],
              ["Responsable LCB-FT", "Pilote le dispositif, rend compte annuellement, valide les cas à risque élevé"],
              ["Déclarant / correspondant Tracfin", "Effectue les DS et répond aux droits de communication"],
              ["Opérationnels (1re ligne)", "KYC, vigilance constante, remontée d'alertes"],
              ["Contrôle permanent / périodique", "Vérifie l'efficacité du dispositif (2e et 3e lignes)"]],
          header=["Acteur", "Rôle"], widths=[5, 12])
    d.add_heading("4. Classification des risques", 1)
    d.add_paragraph("La classification est revue au moins annuellement et à chaque évolution significative (nouveau produit, pays, canal). "
                    "Principaux facteurs identifiés :")
    sub = factors[factors["secteur"] == sector_code].sort_values("risque_brut", ascending=False).head(10)
    table(d, [[r.categorie, r.facteur, r.risque_brut] for r in sub.itertuples()], header=["Catégorie", "Facteur", "Risque brut (/16)"], widths=[3, 11, 3])
    d.add_heading("5. Vigilance à l'égard de la clientèle", 1)
    for t in ["Identification et vérification de l'identité du client et du bénéficiaire effectif avant l'entrée en relation.",
              "Recueil de l'objet et de la nature de la relation, des revenus/patrimoine ou du chiffre d'affaires.",
              "Vigilance simplifiée, standard ou renforcée selon le profil (PPE, pays tiers à haut risque, opérations atypiques).",
              "Mise à jour périodique : risque élevé 1 an, standard 3 ans, faible 5 ans (à adapter).",
              "Remédiation du stock clients fondée sur les risques (AMLR art. 26)."]:
        d.add_paragraph(t, style="List Bullet")
    d.add_heading("6. Surveillance des opérations", 1)
    d.add_paragraph("Les scénarios d'alerte intègrent la situation financière déclarée et le fonctionnement prévisionnel du compte ; les liens avec "
                    "les pays à risque sont détectés sur les opérations effectives ; les stocks et délais de traitement sont pilotés (enseignements "
                    "de la décision ACPR Banque Delubac du 19/06/2025).")
    d.add_heading("7. Typologies de référence", 1)
    for r in typologies[typologies["secteurs"].str.contains(sector_code)].head(8).itertuples():
        d.add_paragraph(f"{r.intitule} — {r.criteres_alerte}", style="List Bullet")
    d.add_heading("8. Déclaration de soupçon et droit d'opposition", 1)
    d.add_paragraph("Toute opération dont l'organisme sait, soupçonne ou a de bonnes raisons de soupçonner qu'elle provient d'une infraction punie "
                    "d'une peine privative de liberté supérieure à un an ou participe au financement du terrorisme fait l'objet d'une DS (CMF L.561-15).")
    d.add_heading("9. Gel des avoirs et sanctions", 1)
    d.add_paragraph("Filtrage des clients, bénéficiaires effectifs et contreparties contre les listes UE, ONU et nationale (DG Trésor) à l'entrée en "
                    "relation, à chaque mise à jour des listes et, pour les paiements instantanés, au moins quotidiennement (Règl. (UE) 2024/886).")
    d.add_heading("10. Conservation, formation, contrôle", 1)
    d.add_paragraph("Conservation 5 ans ; formation annuelle obligatoire avec attestation ; plan de contrôle permanent et périodique.")
    signatures(d, ["Responsable LCB-FT", "Direction générale"])
    return to_bytes(d)


def acceptance_policy(company: str, sector_label: str, countries: pd.DataFrame) -> bytes:
    d = new_doc("Politique d'acceptation clients", company, f"Secteur : {sector_label}")
    d.add_heading("1. Clients interdits", 1)
    for t in ["Personnes et entités inscrites sur les listes de gel UE / ONU / France", "Banques fictives (shell banks)",
              "Clients refusant de fournir les justificatifs requis", "Résidents de pays sous embargo global"]:
        d.add_paragraph(t, style="List Bullet")
    emb = countries[countries["niveau_sanctions"] >= 3]["pays"].tolist()
    d.add_paragraph("Pays sous embargo global (référentiel actuel) : " + ", ".join(emb))
    d.add_heading("2. Clients soumis à autorisation préalable (vigilance renforcée)", 1)
    hr = countries[countries["niveau_risque_pays"].isin(["Très élevé"]) | (countries["gafi_statut"] != "-")]["pays"].tolist()
    for t in ["PPE, membres de leur famille et proches associés", "Clients liés à un pays tiers à haut risque (liste UE / GAFI)",
              "Structures complexes (trusts, fiducies, > 3 niveaux de détention)", "Activités à forte intensité d'espèces", "Clients avec presse négative"]:
        d.add_paragraph(t, style="List Bullet")
    d.add_paragraph("Pays nécessitant une autorisation préalable : " + ", ".join(hr))
    d.add_heading("3. Matrice de décision", 1)
    table(d, [["Faible", "Chargé de clientèle", "5 ans"], ["Standard", "Chargé + contrôle KYC", "3 ans"],
              ["Élevé", "Responsable LCB-FT", "1 an"], ["Très élevé / PPE", "Direction + responsable LCB-FT", "1 an + surveillance renforcée"]],
          header=["Niveau de risque", "Validation", "Revue"], widths=[4, 7, 6])
    signatures(d, ["Responsable LCB-FT", "Direction générale"])
    return to_bytes(d)


def remediation_letter(company: str, client: str, missing: list[str]) -> bytes:
    d = new_doc("Demande de mise à jour de votre dossier", company, "Obligations réglementaires de connaissance client (CMF L.561-6)", confid="Courrier client")
    d.add_paragraph(f"Madame, Monsieur,\n\nDans le cadre de nos obligations légales de vigilance, nous vous remercions de bien vouloir nous transmettre, "
                    f"dans un délai de 30 jours, les éléments suivants concernant le dossier « {client} » :")
    for m in missing:
        d.add_paragraph(m, style="List Bullet")
    d.add_paragraph("À défaut de réponse, nous pourrions être contraints de restreindre certains services, conformément à la réglementation. "
                    "Nous restons à votre disposition pour toute question.\n\nVeuillez agréer, Madame, Monsieur, nos salutations distinguées.")
    d.add_paragraph("Le service Conformité")
    return to_bytes(d)


def training_attendance(company: str, module: str) -> bytes:
    d = new_doc("Feuille d'émargement — formation LCB-FT", company, f"Module : {module} — CMF L.561-34")
    fields(d, ["Date", "Durée", "Formateur", "Modalité (présentiel / e-learning)"])
    table(d, [["", "", "", ""] for _ in range(20)], header=["Nom", "Service", "Signature", "Score quiz"], widths=[5, 4, 5, 3])
    return to_bytes(d)


def gap_report(company: str, sector_label: str, results: pd.DataFrame, score: float) -> bytes:
    d = new_doc("Rapport de mise en conformité", company, f"Secteur : {sector_label} — score de maturité {score:.0f} %")
    d.add_heading("Synthèse", 1)
    d.add_paragraph(f"Le diagnostic couvre {len(results)} exigences. Score global de maturité : {score:.0f} %. Les écarts ci-dessous sont classés "
                    "par priorité et assortis d'une action recommandée.")
    gaps = results[results["statut"] != "Conforme"]
    table(d, [[r.domaine, r.exigence, r.statut, r.reference] for r in gaps.itertuples()], header=["Domaine", "Exigence", "Statut", "Référence"],
          widths=[3, 8, 2.5, 3.5])
    d.add_heading("Points conformes", 1)
    for r in results[results["statut"] == "Conforme"].itertuples():
        d.add_paragraph(r.exigence, style="List Bullet")
    signatures(d, ["Consultant / Responsable conformité", "Direction"])
    return to_bytes(d)


def adverse_report(company: str, subject: str, items: pd.DataFrame, score: float) -> bytes:
    d = new_doc("Rapport de recherche presse négative", company, f"Sujet : {subject} — score {score:.0f}/100")
    d.add_paragraph(f"Recherche réalisée le {date.today():%d/%m/%Y} en sources ouvertes (actualités FR/EN). Résultats à qualifier par l'analyste.")
    if len(items):
        table(d, [[r.published, r.title, r.source, r.mots_cles, r.score] for r in items.itertuples()],
              header=["Date", "Titre", "Source", "Mots-clés", "Score"], widths=[2.2, 8, 2.8, 2.5, 1.5])
    else:
        d.add_paragraph("Aucun article pertinent identifié.")
    fields(d, ["Conclusion de l'analyste", "Impact sur le niveau de risque"])
    signatures(d, ["Analyste", "Responsable LCB-FT"])
    return to_bytes(d)
